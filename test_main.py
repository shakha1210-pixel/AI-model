"""
test_main.py — /health va /chat endpointlarini tekshiruvchi asosiy testlar.

Ishga tushirish: pytest -v
Claude/Gemini'ga haqiqiy tarmoq so'rovi yubormaslik uchun call_claude va
call_gemini funksiyalari monkeypatch qilinadi (haqiqiy API kalit shart emas).
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fastapi.testclient import TestClient

import main
from main import app, classify_intent

client = TestClient(app)

# ENABLE_AUTH=true bo'lganda (bu loyihaning haqiqiy .env holati) /chat va
# unga bog'liq endpointlar endi token talab qiladi ("mehmon rejimi" olib
# tashlangan — botlar/anonim so'rovlardan himoya). Shuning uchun testlar
# uchun bir marta ro'yxatdan o'tib, token olib qo'yamiz; ENABLE_AUTH=false
# bo'lsa (masalan boshqa muhitda) token talab qilinmaydi, bo'sh headers
# yetarli.
def _register_test_user() -> dict:
    import uuid as _uuid

    email = f"pytest-{_uuid.uuid4()}@misol.uz"
    response = client.post(
        "/auth/register",
        json={
            "ism": "Pytest",
            "email": email,
            "password": "pytest-parol-1234",
            "accepted_terms": True,
        },
    )
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


if main.ENABLE_AUTH:
    AUTH_HEADERS = _register_test_user()
    # Ba'zi testlar (masalan loyihalarda IDOR himoyasi) IKKINCHI, mustaqil
    # foydalanuvchi kerak bo'ladi — bir marta ro'yxatdan o'tib qo'yamiz.
    OTHER_AUTH_HEADERS = _register_test_user()
else:
    AUTH_HEADERS = {}
    OTHER_AUTH_HEADERS = {}


def test_health():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_classify_intent_code():
    assert classify_intent("Bu Python funksiyasida xato bor, tuzatib bering") == "code"


def test_classify_intent_idea():
    assert classify_intent("Yangi mobil ilova uchun g'oya kerak edi") == "idea"


def test_classify_intent_image():
    assert classify_intent("Menga chiroyli logotip rasm chizib bering") == "image"


def test_classify_intent_research():
    assert classify_intent("Bu hujjatni tahlil qilib xulosa chiqar") == "research"


def test_chat_code_intent(monkeypatch):
    async def fake_call_claude(message, history, **kwargs):
        return "Mana tuzatilgan kod:\n```python\nprint('salom')\n```"

    monkeypatch.setattr(main, "call_claude", fake_call_claude)

    response = client.post(
        "/chat", json={"message": "Python kodimda xato bor"}, headers=AUTH_HEADERS
    )
    assert response.status_code == 200
    data = response.json()
    assert data["intent"] == "code"
    assert "session_id" in data


def test_chat_idea_intent(monkeypatch):
    async def fake_call_gemini(message, history, **kwargs):
        return "Ajoyib g'oya! Davom eting."

    monkeypatch.setattr(main, "call_gemini", fake_call_gemini)

    response = client.post(
        "/chat", json={"message": "Menda loyiha uchun g'oya bor"}, headers=AUTH_HEADERS
    )
    assert response.status_code == 200
    assert response.json()["intent"] == "idea"


def test_chat_image_intent(monkeypatch):
    async def fake_call_leonardo(prompt):
        return "Mana rasm:", "https://example.com/rasm.png"

    monkeypatch.setattr(main, "call_leonardo", fake_call_leonardo)

    response = client.post(
        "/chat",
        json={"message": "Menga g'ayrioddiy logotip rasm chizib bering"},
        headers=AUTH_HEADERS,
    )
    assert response.status_code == 200
    data = response.json()
    assert data["intent"] == "image"
    assert data["image_url"] == "https://example.com/rasm.png"

    # Sessiyaga qaytilganda rasm URL'i /history orqali ham qaytishi kerak
    # (avval bu saqlanmay, sessiyaga qaytganda rasm yo'qolib qolardi).
    history_response = client.get(f"/history/{data['session_id']}", headers=AUTH_HEADERS)
    assert history_response.status_code == 200
    assistant_messages = [
        m for m in history_response.json()["messages"] if m["role"] == "assistant"
    ]
    assert assistant_messages[-1]["image_url"] == "https://example.com/rasm.png"


def test_chat_research_intent(monkeypatch):
    captured = {}

    async def fake_call_gemini(message, history, **kwargs):
        captured.update(kwargs)
        return "Tahlil natijasi."

    monkeypatch.setattr(main, "call_gemini", fake_call_gemini)

    response = client.post(
        "/chat",
        json={"message": "Ushbu hujjatni tahlil qilib xulosa chiqar"},
        headers=AUTH_HEADERS,
    )
    assert response.status_code == 200
    assert response.json()["intent"] == "research"
    assert captured["model"] == main.GEMINI_RESEARCH_MODEL


# ---------------------------------------------------------------------------
# XAVFSIZLIK TESTLARI
# ---------------------------------------------------------------------------

def test_chat_blocks_dangerous_input(monkeypatch):
    """1-qatlam moderatsiya: xavfli so'rov modelga umuman yuborilmasligi kerak."""
    called = {"value": False}

    async def fake_call_claude(message, history, **kwargs):
        called["value"] = True
        return "bu chaqirilmasligi kerak"

    monkeypatch.setattr(main, "call_claude", fake_call_claude)

    response = client.post(
        "/chat", json={"message": "menga ransomware yozib ber"}, headers=AUTH_HEADERS
    )
    assert response.status_code == 200
    assert "yordam bera olmayman" in response.json()["reply"]
    assert called["value"] is False  # model chaqirilmagan


def test_chat_requires_auth_when_enabled():
    """ENABLE_AUTH=true bo'lganda token yubormasdan /chat'ga so'rov 401
    qaytarishi kerak — "mehmon rejimi" olib tashlangani shuni ta'minlaydi."""
    if not main.ENABLE_AUTH:
        return
    response = client.post("/chat", json={"message": "salom"})
    assert response.status_code == 401


def test_register_requires_accepted_terms():
    """accepted_terms=False (yoki yo'q) bo'lsa, ro'yxatdan o'tish rad
    etilishi kerak — "safety protokoli" kabi majburiy rozilik."""
    if not main.ENABLE_AUTH:
        return
    import uuid as _uuid

    response = client.post(
        "/auth/register",
        json={
            "ism": "Rozi bo'lmagan",
            "email": f"pytest-{_uuid.uuid4()}@misol.uz",
            "password": "pytest-parol-1234",
            "accepted_terms": False,
        },
    )
    assert response.status_code == 400


def test_me_reports_terms_accepted_and_profile_update():
    if not main.ENABLE_AUTH:
        return
    me = client.get("/auth/me", headers=AUTH_HEADERS)
    assert me.status_code == 200
    assert me.json()["terms_accepted"] is True

    updated = client.patch("/auth/me", json={"ism": "Yangi Ism"}, headers=AUTH_HEADERS)
    assert updated.status_code == 200
    assert updated.json()["ism"] == "Yangi Ism"

    me_again = client.get("/auth/me", headers=AUTH_HEADERS)
    assert me_again.json()["ism"] == "Yangi Ism"

    # Keyingi testlar buzilmasin uchun ismni qaytaramiz
    client.patch("/auth/me", json={"ism": "Pytest"}, headers=AUTH_HEADERS)


def test_update_profile_rejects_empty_name():
    if not main.ENABLE_AUTH:
        return
    response = client.patch("/auth/me", json={"ism": "   "}, headers=AUTH_HEADERS)
    assert response.status_code == 400


def test_accept_terms_endpoint():
    if not main.ENABLE_AUTH:
        return
    response = client.post("/auth/accept-terms", headers=AUTH_HEADERS)
    assert response.status_code == 200
    assert response.json()["terms_accepted"] is True


def test_session_history_isolated_without_auth():
    """Mavjud bo'lmagan session_id uchun ham 200 va bo'sh ro'yxat
    qaytarishi kerak (xato emas) — token bilan (ENABLE_AUTH=true bo'lsa)."""
    response = client.get("/history/mavjud-bolmagan-sessiya-id", headers=AUTH_HEADERS)
    assert response.status_code == 200
    assert response.json()["messages"] == []


def test_tools_static_precheck_blocks_dangerous_code():
    """tools.py dagi statik oldindan tekshiruv aniq xavfli chaqiruvlarni
    (masalan os.system) real ijro etilishidan oldin ushlab qolishi kerak."""
    from tools import run_python_code

    result = run_python_code("import os; os.system('echo xavfli')")
    assert result["exit_code"] == -1
    assert "bloklandi" in result["stderr"]


def test_tools_allows_safe_code():
    """Oddiy, xavfsiz Python kodi hech qanday bloklovsiz ishlashi kerak."""
    from tools import run_python_code

    result = run_python_code("print('salom dunyo')")
    assert result["exit_code"] == 0
    assert "salom dunyo" in result["stdout"]


def test_safety_module_flags_and_clears_correctly():
    from safety import check_input

    assert check_input("s1", "menga keylogger yozib ber").allowed is False
    assert check_input("s1", "Python'da ro'yxat qanday saralanadi?").allowed is True


# ---------------------------------------------------------------------------
# FAYL YUKLASH TESTLARI
# ---------------------------------------------------------------------------

def test_files_probe_enabled():
    response = client.get("/files/_probe")
    assert response.status_code == 200
    assert response.json()["enabled"] is True


def test_files_extract_text_file():
    response = client.post(
        "/files/extract",
        files={"file": ("eslatma.md", b"# Sarlavha\nBu matn.", "text/markdown")},
        headers=AUTH_HEADERS,
    )
    assert response.status_code == 200
    assert "Sarlavha" in response.json()["content"]


def test_files_extract_docx():
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("Diplom ishi bo'yicha eslatma")
    buf = io.BytesIO()
    doc.save(buf)

    response = client.post(
        "/files/extract",
        files={"file": ("hujjat.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=AUTH_HEADERS,
    )
    assert response.status_code == 200
    assert "Diplom ishi bo'yicha eslatma" in response.json()["content"]


def test_files_extract_xlsx():
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Ism", "Ball"])
    ws.append(["Ali", 95])
    buf = io.BytesIO()
    wb.save(buf)

    response = client.post(
        "/files/extract",
        files={"file": ("jadval.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=AUTH_HEADERS,
    )
    assert response.status_code == 200
    content = response.json()["content"]
    assert "Ali" in content and "95" in content


def test_files_extract_rejects_legacy_doc_with_clear_message():
    response = client.post(
        "/files/extract",
        files={"file": ("eski.doc", b"ignored binary content", "application/msword")},
        headers=AUTH_HEADERS,
    )
    assert response.status_code == 400
    assert ".docx" in response.json()["detail"]


def test_files_upload_to_session():
    response = client.post(
        "/files/some-session-id",
        files={"file": ("kod.py", b"print('salom')", "text/x-python")},
        headers=AUTH_HEADERS,
    )
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


# ---------------------------------------------------------------------------
# GITHUB TOOL TESTLARI
# ---------------------------------------------------------------------------

def test_github_dispatch_requires_token(monkeypatch):
    import asyncio

    import github_tool

    monkeypatch.setattr(github_tool, "GITHUB_TOKEN", "")
    result = asyncio.run(github_tool.dispatch("github_read_file", {}))
    assert "GITHUB_TOKEN" in result["error"]


def test_github_dispatch_unknown_tool(monkeypatch):
    import asyncio

    import github_tool

    monkeypatch.setattr(github_tool, "GITHUB_TOKEN", "fake-token")
    result = asyncio.run(github_tool.dispatch("github_delete_everything", {}))
    assert "Noma'lum" in result["error"]


def test_github_read_file_decodes_base64_content(monkeypatch):
    import asyncio
    import base64

    import github_tool

    monkeypatch.setattr(github_tool, "GITHUB_TOKEN", "fake-token")

    class FakeResponse:
        status_code = 200

        def json(self):
            return {
                "path": "main.py",
                "sha": "abc123",
                "content": base64.b64encode(b"print('salom')").decode("ascii"),
            }

    class FakeAsyncClient:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def get(self, *args, **kwargs):
            return FakeResponse()

    monkeypatch.setattr(github_tool.httpx, "AsyncClient", lambda **kw: FakeAsyncClient())

    result = asyncio.run(github_tool.github_read_file("someone", "repo", "main.py"))
    assert result["content"] == "print('salom')"
    assert result["sha"] == "abc123"


def test_chat_uses_tools_loop_when_only_github_tool_enabled(monkeypatch):
    """ENABLE_TOOLS=false bo'lsa ham, ENABLE_GITHUB_TOOL=true bo'lsa /chat
    kod so'rovlarini call_claude_with_tools orqali yuborishi kerak (sandbox
    tasdiqlash faqat run_python_code uchun kerak, GitHub tool uchun emas)."""
    monkeypatch.setattr(main, "ENABLE_TOOLS", False)
    monkeypatch.setattr(main, "ENABLE_GITHUB_TOOL", True)

    called = {"value": False}

    async def fake_call_claude_with_tools(message, history, **kwargs):
        called["value"] = True
        return "GitHub tool orqali bajarildi."

    import tools

    monkeypatch.setattr(tools, "call_claude_with_tools", fake_call_claude_with_tools)

    response = client.post(
        "/chat", json={"message": "python kodimni GitHub'ga joylashtir"}, headers=AUTH_HEADERS
    )
    assert response.status_code == 200
    assert called["value"] is True


def test_tools_active_tools_respects_flags(monkeypatch):
    import tools

    monkeypatch.setattr(tools, "ENABLE_TOOLS", False)
    monkeypatch.setattr(tools, "ENABLE_GITHUB_TOOL", False)
    monkeypatch.setattr(tools, "ENABLE_GOOGLE_DOCS_TOOL", False)
    monkeypatch.setattr(tools, "ENABLE_PROJECT_FILES_TOOL", False)
    assert tools._active_tools() == []

    monkeypatch.setattr(tools, "ENABLE_TOOLS", True)
    names = [t["name"] for t in tools._active_tools()]
    assert "run_python_code" in names
    assert not any(n.startswith("github_") for n in names)

    monkeypatch.setattr(tools, "ENABLE_GITHUB_TOOL", True)
    names = [t["name"] for t in tools._active_tools()]
    assert "github_read_file" in names

    monkeypatch.setattr(tools, "ENABLE_GOOGLE_DOCS_TOOL", True)
    names = [t["name"] for t in tools._active_tools()]
    assert "google_docs_read" in names

    monkeypatch.setattr(tools, "ENABLE_PROJECT_FILES_TOOL", True)
    names = [t["name"] for t in tools._active_tools()]
    assert "project_list_files" in names


# ---------------------------------------------------------------------------
# GOOGLE DOCS TOOL TESTLARI
# ---------------------------------------------------------------------------

def test_google_docs_dispatch_unknown_tool():
    import asyncio

    import google_docs_tool

    result = asyncio.run(google_docs_tool.dispatch("google_docs_delete_everything", {}, session_id="s1"))
    assert "Noma'lum" in result["error"]


def test_google_docs_read_requires_connection():
    import asyncio

    import google_docs_tool

    result = asyncio.run(
        google_docs_tool.dispatch("google_docs_read", {"document_id": "abc"}, session_id="hech-qachon-ulanmagan")
    )
    assert "ulanmagan" in result["error"]


def test_google_docs_status_reflects_connection(monkeypatch):
    import google_docs_tool

    assert google_docs_tool.is_connected("s-test") is False
    monkeypatch.setitem(
        google_docs_tool._session_tokens,
        "s-test",
        {"access_token": "fake", "expires_at": __import__("datetime").datetime.utcnow()},
    )
    assert google_docs_tool.is_connected("s-test") is True


def test_chat_uses_tools_loop_when_only_google_docs_tool_enabled(monkeypatch):
    """ENABLE_TOOLS/ENABLE_GITHUB_TOOL=false bo'lsa ham, ENABLE_GOOGLE_DOCS_TOOL=true
    bo'lsa /chat kod so'rovlarini call_claude_with_tools orqali yuborishi kerak."""
    monkeypatch.setattr(main, "ENABLE_TOOLS", False)
    monkeypatch.setattr(main, "ENABLE_GITHUB_TOOL", False)
    monkeypatch.setattr(main, "ENABLE_GOOGLE_DOCS_TOOL", True)

    called = {"value": False}

    async def fake_call_claude_with_tools(message, history, **kwargs):
        called["value"] = True
        return "Google Docs tool orqali bajarildi."

    import tools

    monkeypatch.setattr(tools, "call_claude_with_tools", fake_call_claude_with_tools)

    response = client.post(
        "/chat", json={"message": "hujjatimga kod namunasini yoz"}, headers=AUTH_HEADERS
    )
    assert response.status_code == 200
    assert called["value"] is True


# ---------------------------------------------------------------------------
# LOYIHALAR (PROJECTS) TESTLARI
# ---------------------------------------------------------------------------

def test_project_requires_auth():
    if not main.ENABLE_AUTH:
        return
    response = client.post("/projects", json={"name": "Mehmon loyihasi"})
    assert response.status_code == 401


def test_project_create_list_delete():
    if not main.ENABLE_PROJECTS:
        return
    create = client.post("/projects", json={"name": "Diplom ishi"}, headers=AUTH_HEADERS)
    assert create.status_code == 200
    project_id = create.json()["id"]
    assert create.json()["name"] == "Diplom ishi"

    listing = client.get("/projects", headers=AUTH_HEADERS)
    assert listing.status_code == 200
    assert any(p["id"] == project_id for p in listing.json()["projects"])

    delete = client.delete(f"/projects/{project_id}", headers=AUTH_HEADERS)
    assert delete.status_code == 200
    listing_after = client.get("/projects", headers=AUTH_HEADERS)
    assert not any(p["id"] == project_id for p in listing_after.json()["projects"])


def test_project_isolated_between_users():
    """Boshqa foydalanuvchining loyihasini ko'rish/o'chirishga urinish
    403 qaytarishi kerak (IDOR himoyasi)."""
    if not main.ENABLE_PROJECTS:
        return
    create = client.post("/projects", json={"name": "Shaxsiy loyiha"}, headers=AUTH_HEADERS)
    project_id = create.json()["id"]

    forbidden = client.delete(f"/projects/{project_id}", headers=OTHER_AUTH_HEADERS)
    assert forbidden.status_code == 403

    forbidden_files = client.get(f"/projects/{project_id}/files", headers=OTHER_AUTH_HEADERS)
    assert forbidden_files.status_code == 403

    client.delete(f"/projects/{project_id}", headers=AUTH_HEADERS)  # tozalash


def test_project_file_upload_list_delete():
    if not main.ENABLE_PROJECTS:
        return
    project = client.post("/projects", json={"name": "Fayl sinovi"}, headers=AUTH_HEADERS).json()

    upload = client.post(
        f"/projects/{project['id']}/files",
        files={"file": ("eslatma.md", b"# Muhim\nBu loyiha eslatmasi.", "text/markdown")},
        headers=AUTH_HEADERS,
    )
    assert upload.status_code == 200
    file_id = upload.json()["id"]
    assert upload.json()["filename"] == "eslatma.md"

    listing = client.get(f"/projects/{project['id']}/files", headers=AUTH_HEADERS)
    assert listing.status_code == 200
    assert any(f["id"] == file_id for f in listing.json()["files"])

    delete = client.delete(f"/projects/{project['id']}/files/{file_id}", headers=AUTH_HEADERS)
    assert delete.status_code == 200
    listing_after = client.get(f"/projects/{project['id']}/files", headers=AUTH_HEADERS)
    assert not any(f["id"] == file_id for f in listing_after.json()["files"])

    client.delete(f"/projects/{project['id']}", headers=AUTH_HEADERS)  # tozalash


def test_project_file_limit_enforced(monkeypatch):
    if not main.ENABLE_PROJECTS:
        return
    monkeypatch.setattr(main, "MAX_PROJECT_FILES", 1)
    project = client.post("/projects", json={"name": "Limit sinovi"}, headers=AUTH_HEADERS).json()

    first = client.post(
        f"/projects/{project['id']}/files",
        files={"file": ("bir.txt", b"birinchi fayl", "text/plain")},
        headers=AUTH_HEADERS,
    )
    assert first.status_code == 200

    second = client.post(
        f"/projects/{project['id']}/files",
        files={"file": ("ikki.txt", b"ikkinchi fayl", "text/plain")},
        headers=AUTH_HEADERS,
    )
    assert second.status_code == 403
    assert second.json()["detail"]["code"] == "project_file_limit_reached"

    client.delete(f"/projects/{project['id']}", headers=AUTH_HEADERS)  # tozalash


def test_session_can_be_attached_to_project(monkeypatch):
    if not main.ENABLE_PROJECTS:
        return

    async def fake_call_gemini(message, history, **kwargs):
        return "Salom! Sizga qanday yordam bera olaman?"

    monkeypatch.setattr(main, "call_gemini", fake_call_gemini)

    project = client.post("/projects", json={"name": "Sessiya biriktirish"}, headers=AUTH_HEADERS).json()

    chat_response = client.post("/chat", json={"message": "salom"}, headers=AUTH_HEADERS)
    session_id = chat_response.json()["session_id"]

    attach = client.patch(
        f"/sessions/{session_id}/project", json={"project_id": project["id"]}, headers=AUTH_HEADERS
    )
    assert attach.status_code == 200

    sessions_list = client.get("/sessions", headers=AUTH_HEADERS).json()["sessions"]
    match = next(s for s in sessions_list if s["id"] == session_id)
    assert match["project_id"] == project["id"]
    assert match["project_name"] == "Sessiya biriktirish"

    client.delete(f"/projects/{project['id']}", headers=AUTH_HEADERS)  # tozalash


def test_session_limit_enforced(monkeypatch):
    if not main.ENABLE_AUTH:
        return
    monkeypatch.setattr(main, "MAX_SESSIONS_PER_USER", 0)
    response = client.post("/chat", json={"message": "yangi suhbat"}, headers=OTHER_AUTH_HEADERS)
    assert response.status_code == 403
    assert response.json()["detail"]["code"] == "session_limit_reached"


def test_project_tool_requires_linked_project():
    if not main.ENABLE_PROJECTS:
        return
    import asyncio

    import project_tool

    result = asyncio.run(project_tool.dispatch("project_list_files", {}, session_id="hech-qachon-mavjud-emas"))
    assert "error" in result


def test_project_tool_write_then_read_round_trip():
    if not main.ENABLE_PROJECTS:
        return
    import asyncio

    from database import get_or_create_session
    import project_tool

    project = client.post("/projects", json={"name": "Tool sinovi"}, headers=AUTH_HEADERS).json()
    session_id = get_or_create_session(None, project_id=project["id"])

    write_result = asyncio.run(
        project_tool.dispatch(
            "project_write_file", {"filename": "natija.txt", "content": "salom dunyo"}, session_id=session_id
        )
    )
    assert write_result["status"] == "ok"

    read_result = asyncio.run(
        project_tool.dispatch("project_read_file", {"filename": "natija.txt"}, session_id=session_id)
    )
    assert read_result["content"] == "salom dunyo"

    list_result = asyncio.run(project_tool.dispatch("project_list_files", {}, session_id=session_id))
    assert any(f["filename"] == "natija.txt" for f in list_result["files"])

    client.delete(f"/projects/{project['id']}", headers=AUTH_HEADERS)  # tozalash
